"""
Leva para o acervo o padrão de cada preparação que a autora marcou na planilha
"Livro Vivo - padrao das preparacoes.xlsx" (aba "Preparações", coluna G).

O assistente lê o campo `padrao` de cada preparação em
public/app/acervo/assistente.json. Preparação sem o campo continua com a
classificação provisória, feita pelo custo dos ingredientes (nivelDoPrato em
public/app/assets/assistente.js).

Uso:
  python scripts/aplicar-padroes.py planilha.xlsx
  python scripts/aplicar-padroes.py planilha.xlsx --linhas "Prato principal,Acompanhamento"

--linhas aplica só as linhas do cardápio que a autora disse ter revisado.
Precisa de openpyxl (pip install openpyxl).

ATENÇÃO: este repositório é público. A planilha tem o acervo inteiro e não
deve ser comitada.
"""
import argparse
import json
import pathlib
import sys

from openpyxl import load_workbook

RAIZ = pathlib.Path(__file__).resolve().parent.parent
ACERVO = RAIZ / 'public' / 'app' / 'acervo' / 'assistente.json'
VALOR = {'popular': 'popular', 'médio': 'medio', 'medio': 'medio', 'luxo': 'luxo'}

ap = argparse.ArgumentParser()
ap.add_argument('planilha')
ap.add_argument('--linhas', help='linhas do cardápio revisadas, separadas por vírgula')
args = ap.parse_args()
so_linhas = {x.strip().lower() for x in args.linhas.split(',')} if args.linhas else None

ws = load_workbook(args.planilha, read_only=True)['Preparações']
marcados, problemas = {}, []
for i, linha in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    pid, linha_cardapio, _, nome, _, _, padrao = linha[:7]
    if pid is None:
        continue
    if so_linhas and str(linha_cardapio).strip().lower() not in so_linhas:
        continue
    valor = VALOR.get(str(padrao or '').strip().lower())
    if not valor:
        problemas.append(f'linha {i} ({nome}): padrão "{padrao}" não reconhecido')
        continue
    marcados[int(pid)] = valor

if problemas:
    print('Nada foi gravado. Corrija na planilha:', *problemas, sep='\n  ')
    sys.exit(1)

acervo = json.loads(ACERVO.read_text(encoding='utf-8'))
mudou = 0
for p in acervo['preparacoes']:
    if p['id'] in marcados and p.get('padrao') != marcados[p['id']]:
        p['padrao'] = marcados[p['id']]
        mudou += 1
faltando = set(marcados) - {p['id'] for p in acervo['preparacoes']}

ACERVO.write_text(json.dumps(acervo, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(f'{len(marcados)} preparações lidas, {mudou} atualizadas no acervo.')
if faltando:
    print(f'{len(faltando)} IDs da planilha não existem no acervo: {sorted(faltando)[:20]}')
